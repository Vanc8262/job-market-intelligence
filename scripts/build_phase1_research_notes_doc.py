"""Build the updated Phase 1 research notes DOCX."""
# ruff: noqa: ANN001, E501

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

PROJECT_ROOT = Path(__file__).resolve().parents[1]
OUTPUT_PATH = (
    PROJECT_ROOT / "docs" / "Job Market Intelligence Research Notes - Phase 1 Updated.docx"
)

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(31, 41, 55)
MUTED = RGBColor(88, 96, 105)
LIGHT_FILL = "F2F4F7"
CALLOUT_FILL = "F4F6F9"
BORDER = "D9DEE7"


def set_cell_shading(cell, fill: str) -> None:
    """Apply a background fill to a table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(
    cell, top: int = 80, bottom: int = 80, start: int = 120, end: int = 120
) -> None:
    """Set cell margins in DXA."""
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin_name, value in {
        "top": top,
        "bottom": bottom,
        "start": start,
        "end": end,
    }.items():
        node = tc_mar.find(qn(f"w:{margin_name}"))
        if node is None:
            node = OxmlElement(f"w:{margin_name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table) -> None:
    """Apply a quiet single-grid table border."""
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), BORDER)


def set_table_geometry(table, widths: list[float]) -> None:
    """Set fixed table geometry for all rows."""
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    table.allow_autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), "9360")
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    for row in table.rows:
        for index, width in enumerate(widths):
            if index >= len(row.cells):
                continue
            row.cells[index].width = Inches(width)
            row.cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(row.cells[index])
    set_table_borders(table)


def set_paragraph_spacing(
    paragraph, before: int = 0, after: int = 6, line_spacing: float = 1.1
) -> None:
    """Apply standard paragraph rhythm."""
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line_spacing


def add_run(paragraph, text: str, bold: bool = False, color: RGBColor | None = None) -> None:
    """Append styled text to a paragraph."""
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Calibri"
    run.font.size = Pt(11)
    run.font.color.rgb = color or INK


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    """Add a preset-compliant heading."""
    paragraph = doc.add_heading(text, level=level)
    for run in paragraph.runs:
        run.font.name = "Calibri"
        run.font.color.rgb = BLUE if level < 3 else DARK_BLUE
        run.font.size = Pt(16 if level == 1 else 13 if level == 2 else 12)
    paragraph.paragraph_format.space_before = Pt(16 if level == 1 else 12 if level == 2 else 8)
    paragraph.paragraph_format.space_after = Pt(8 if level == 1 else 6 if level == 2 else 4)


def add_body(doc: Document, text: str) -> None:
    """Add body text."""
    paragraph = doc.add_paragraph()
    add_run(paragraph, text)
    set_paragraph_spacing(paragraph)


def add_bullet(doc: Document, text: str) -> None:
    """Add a bulleted paragraph."""
    paragraph = doc.add_paragraph(style="List Bullet")
    add_run(paragraph, text)
    set_paragraph_spacing(paragraph, after=4, line_spacing=1.167)


def add_callout(doc: Document, label: str, text: str) -> None:
    """Add a one-column callout table."""
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [6.5])
    cell = table.cell(0, 0)
    set_cell_shading(cell, CALLOUT_FILL)
    paragraph = cell.paragraphs[0]
    add_run(paragraph, f"{label}: ", bold=True, color=DARK_BLUE)
    add_run(paragraph, text)


def add_table(
    doc: Document, headers: list[str], rows: list[list[str]], widths: list[float]
) -> None:
    """Add a fixed-width table with a shaded header row."""
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths)
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, LIGHT_FILL)
        paragraph = cell.paragraphs[0]
        add_run(paragraph, header, bold=True, color=DARK_BLUE)
    for row_values in rows:
        row = table.add_row()
        for idx, value in enumerate(row_values):
            cell = row.cells[idx]
            set_cell_margins(cell)
            paragraph = cell.paragraphs[0]
            add_run(paragraph, value)
    set_table_geometry(table, widths)
    doc.add_paragraph()


def configure_document(doc: Document) -> None:
    """Configure page setup and core styles."""
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    for style_name in ("List Bullet", "List Number"):
        style = styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.167


def add_header_footer(doc: Document) -> None:
    """Add understated running document furniture."""
    section = doc.sections[0]
    header = section.header.paragraphs[0]
    header.text = "Job Market Intelligence Platform | Phase 1 Research Notes"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.name = "Calibri"
        run.font.size = Pt(9)
        run.font.color.rgb = MUTED

    footer = section.footer.paragraphs[0]
    footer.text = "Prepared from the repository implementation state as of August 17, 2026"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in footer.runs:
        run.font.name = "Calibri"
        run.font.size = Pt(9)
        run.font.color.rgb = MUTED


def build_document() -> None:
    """Build and save the updated research notes document."""
    doc = Document()
    configure_document(doc)
    add_header_footer(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = title.add_run("Job Market Intelligence Research Notes")
    run.font.name = "Calibri"
    run.font.size = Pt(22)
    run.font.bold = True
    run.font.color.rgb = DARK_BLUE
    title.paragraph_format.space_after = Pt(3)

    subtitle = doc.add_paragraph()
    add_run(
        subtitle,
        "Phase 1 implementation update: project foundation, canonical schema, extractor contract, "
        "storage layers, configuration, logging, checkpoints, retry behavior, hashing, and tests.",
        color=MUTED,
    )
    subtitle.paragraph_format.space_after = Pt(12)

    add_callout(
        doc,
        "Phase 1 status",
        "Completed. The repository now contains a tested, source-agnostic foundation for later "
        "data acquisition, processing, enrichment, analytics, and dashboard phases. Phase 1 does "
        "not yet implement live scraping, NLP enrichment, production scheduling, or dashboards.",
    )

    add_heading(doc, "1. Project Vision", 1)
    add_body(
        doc,
        "The Job Market Intelligence Platform (JMIP) is a production-oriented Python foundation "
        "for collecting, preserving, standardizing, enriching, and analyzing technical job postings. "
        "The project targets data science, artificial intelligence, machine learning, data engineering, "
        "software engineering, and adjacent computer science labor-market signals.",
    )
    add_body(
        doc,
        "The long-term research purpose remains labor-market research, career planning, skill-gap "
        "analysis, salary analysis, technology trend forecasting, academic research, and portfolio "
        "demonstration.",
    )

    add_heading(doc, "2. Current Repository Structure", 1)
    add_table(
        doc,
        ["Directory", "Phase 1 role", "Implementation status"],
        [
            [
                "config/",
                "Validated YAML settings for paths, sources, skills, locations, titles, salary, and logging.",
                "Implemented",
            ],
            [
                "models/",
                "Canonical Pydantic job posting schema and related enums/submodels.",
                "Implemented",
            ],
            [
                "extractors/",
                "Shared source-agnostic extractor interface and orchestration result types.",
                "Implemented",
            ],
            [
                "storage/",
                "Bronze, silver, and gold storage abstractions plus file-backed JSON implementations.",
                "Implemented",
            ],
            [
                "utils/",
                "Config loading, logging, checkpointing, retry behavior, hashing, and helpers.",
                "Implemented",
            ],
            [
                "processors/",
                "Validator and normalizer shells for future silver processing.",
                "Foundation placeholders",
            ],
            [
                "analytics/, enrichment/, dashboard/",
                "Future analytics, NLP/AI enrichment, and UI surfaces.",
                "Deferred to later phases",
            ],
            ["scripts/", "Pipeline entry point for foundation validation.", "Implemented"],
            ["tests/", "Unit tests covering Phase 1 infrastructure behavior.", "Implemented"],
        ],
        [1.15, 4.0, 1.35],
    )

    add_heading(doc, "3. Phase 1 Accomplishments", 1)
    add_table(
        doc,
        ["Step", "Completed work", "Evidence in code"],
        [
            [
                "Project structure",
                "Created a modular package layout aligned to the planned pipeline layers.",
                "README.md, pyproject.toml, config/, models/, extractors/, storage/, utils/, tests/",
            ],
            [
                "Configuration",
                "Added a singleton Config loader with Pydantic validation, YAML loading, .env support, and selected JMIP_* overrides.",
                "utils/config.py and config/*.yaml",
            ],
            [
                "Canonical data model",
                "Defined a validated JobPosting schema with source, company, position, location, salary, employment, experience, education, description, skills, raw, and metadata sections.",
                "models/job_posting.py",
            ],
            [
                "Extractor interface",
                "Implemented BaseExtractor with search/fetch/parse hooks, extract() orchestration, raw preservation, checkpointing, result summaries, and source config access.",
                "extractors/base_extractor.py",
            ],
            [
                "Data layers",
                "Created safe local storage contracts and JSON-backed bronze, silver, and gold layers.",
                "storage/base.py, bronze.py, silver.py, gold.py",
            ],
            [
                "Logging",
                "Centralized Loguru setup with console/file sinks, default context, bind_context(), log_event(), and SCRAPER level.",
                "utils/logger.py and config/logging.yaml",
            ],
            [
                "Operational utilities",
                "Added checkpoint save/load/clear, exponential retry decorator, and deterministic SHA-256 identity/content/duplicate hashing.",
                "utils/checkpoint.py, retry.py, hashing.py",
            ],
            [
                "Verification",
                "Added foundation tests for config, model validation, extractor workflow, storage layers, logging context, checkpoints, and hashes.",
                "tests/test_foundation.py",
            ],
        ],
        [1.15, 3.5, 1.85],
    )

    add_heading(doc, "4. Canonical Data Model", 1)
    add_body(
        doc,
        "The canonical record emitted by future extractors is JobPosting. It is intentionally strict "
        "with extra='forbid' on canonical sections, while RawSourceRecord allows source-specific "
        "fields for auditability. The model uses Pydantic validation and assignment validation.",
    )
    add_table(
        doc,
        ["Model section", "Fields and behavior"],
        [
            [
                "PostingMetadata",
                "job_id, content_hash, duplicate_hash, ingestion_batch_id, first_seen_at, last_seen_at, posted_at, expires_at, status, is_active, tags; validates date ordering and status-active consistency.",
            ],
            [
                "SourceInfo",
                "name, source_type, url, external_id, fetched_at; source_type uses SourceType enum.",
            ],
            ["CompanyInfo", "name, normalized_name, website, industry."],
            ["PositionInfo", "title, normalized_title, family, seniority."],
            [
                "LocationInfo",
                "raw, city, region, country, workplace_type, latitude, longitude; validates latitude and longitude bounds.",
            ],
            [
                "SalaryInfo",
                "min_amount, max_amount, currency, pay_period, raw_text; validates ranges and normalizes currency to uppercase.",
            ],
            ["ExperienceInfo", "level, min_years, max_years, raw_text; validates year range."],
            ["EducationInfo", "minimum_degree, preferred_degrees, certifications, raw_text."],
            [
                "DescriptionInfo",
                "raw_text, clean_text, html, responsibilities, qualifications, benefits.",
            ],
            ["SkillInfo", "explicit, inferred, categories."],
            [
                "RawSourceRecord",
                "raw_url, raw_payload, raw_text, plus allowed source-specific extras.",
            ],
        ],
        [1.8, 4.7],
    )
    add_bullet(
        doc,
        "Enums now include SourceType, PostingStatus, EmploymentType, WorkplaceType, PayPeriod, ExperienceLevel, and DegreeLevel.",
    )
    add_bullet(
        doc,
        "JobPosting.with_derived_metadata() fills stable job_id, content_hash, and duplicate_hash from canonical identity fields.",
    )
    add_bullet(
        doc,
        "The model supports analytics questions around active postings, compensation, geography, skills, experience inflation, education requirements, and source provenance.",
    )

    add_heading(doc, "5. Configuration Inventory", 1)
    add_table(
        doc,
        ["Config file", "Current contents"],
        [
            [
                "settings.yaml",
                "Environment, project metadata, data/log paths, storage defaults, scraping retry/checkpoint/timeouts/user agent, and processing batch/strictness settings.",
            ],
            [
                "sources.yaml",
                "Enabled primary ATS sources: Greenhouse, Lever, Ashby. Deferred sources include Indeed, LinkedIn, Workday, and SmartRecruiters with base URLs and rate limits.",
            ],
            [
                "skills.yaml",
                "Skill taxonomy categories: programming languages, data science, machine learning, data engineering, cloud, and AI.",
            ],
            [
                "locations.yaml",
                "Default country, remote/hybrid keywords, regions, and major markets.",
            ],
            [
                "salary.yaml",
                "Supported currencies, pay periods, annual-hours normalization, default pay period, and rough salary bands.",
            ],
            ["titles.yaml", "Role families and seniority keywords for later title normalization."],
            [
                "logging.yaml",
                "Console/file sinks, retention/rotation/compression, and formats that include component, source, job_id, batch, run, and event context.",
            ],
        ],
        [1.45, 5.05],
    )

    add_heading(doc, "6. Extractor Contract", 1)
    add_body(
        doc,
        "BaseExtractor is the shared source interface. Source-specific extractors for Greenhouse, Lever, "
        "Ashby, Workday, Indeed, LinkedIn, or SmartRecruiters should implement only the source-specific "
        "search_jobs(), fetch_job(), and parse_job() methods while reusing the shared extract() workflow.",
    )
    add_bullet(doc, "SearchRequest captures query, optional location, and filters.")
    add_bullet(
        doc,
        "ExtractionResult reports references found, parsed postings, raw records saved, succeeded count, failed count, and failed reference messages.",
    )
    add_bullet(
        doc, "Raw payloads can be persisted through bronze storage before parsing into JobPosting."
    )
    add_bullet(
        doc,
        "Checkpoint hooks support resumable future crawling using source-specific checkpoint names.",
    )

    add_heading(doc, "7. Data Layers", 1)
    add_table(
        doc,
        ["Layer", "Implemented behavior", "Intended use"],
        [
            [
                "Bronze",
                "File-backed JSON storage that preserves raw source records immutably; duplicate keys raise FileExistsError.",
                "Audit trail for HTML, API responses, source JSON, and original payloads.",
            ],
            [
                "Silver",
                "File-backed JSON storage that validates payloads as JobPosting on save and load.",
                "Canonical normalized records for downstream enrichment and deduplication.",
            ],
            [
                "Gold",
                "File-backed JSON analytics payload storage with dataframe-like object normalization.",
                "Analytics-ready tables such as monthly skill counts, salary summaries, hiring velocity, and trend outputs.",
            ],
            [
                "BaseStorage",
                "Safe key-to-path translation with traversal protection and a common save/load/exists/delete contract.",
                "Consistent behavior across future Parquet, DuckDB, or object-store backends.",
            ],
        ],
        [1.0, 3.35, 2.15],
    )

    add_heading(doc, "8. Logging and Operations", 1)
    add_body(
        doc,
        "The logging system is centralized around Loguru and configured through YAML. The implementation "
        "adds stable structured fields so future extraction, processing, and analytics events can be "
        "correlated by component, source, job_id, ingestion_batch_id, run_id, and event.",
    )
    add_bullet(doc, "configure_logging() installs console, application, and error sinks.")
    add_bullet(doc, "LogContext defines standard fields with safe defaults.")
    add_bullet(
        doc, "bind_context() returns a logger pre-bound for a module, source, job, batch, or run."
    )
    add_bullet(
        doc,
        "log_event() records structured events at normal Loguru levels or the custom SCRAPER level.",
    )

    add_heading(doc, "9. Verification Snapshot", 1)
    add_table(
        doc,
        ["Gate", "Current result"],
        [
            ["pytest", "11 tests passed."],
            ["ruff", "All checked files passed lint."],
            ["mypy", "No issues found across 19 source files in the core Phase 1 surface."],
            [
                "Coverage focus",
                "Config loading, Pydantic model validation, extractor workflow, bronze/silver/gold storage, unsafe key rejection, checkpoint round trip, deterministic hashes, and logging context.",
            ],
        ],
        [1.45, 5.05],
    )

    add_heading(doc, "10. Current Boundaries and Phase 2 Handoff", 1)
    add_body(
        doc,
        "Phase 1 intentionally stops at the foundation. The project is now ready for data acquisition "
        "work, but it does not yet perform live crawling, browser automation, NLP skill extraction, "
        "salary parsing, deduplication, DuckDB marts, scheduled jobs, or dashboard analytics.",
    )
    add_table(
        doc,
        ["Next phase priority", "Recommended first implementation step"],
        [
            [
                "Greenhouse extractor",
                "Implement a concrete extractor using the shared BaseExtractor workflow and persist raw API responses into BronzeStorage.",
            ],
            [
                "Lever and Ashby extractors",
                "Follow the same pattern after Greenhouse proves the interface.",
            ],
            [
                "Silver processing",
                "Extend JobPostingNormalizer with configured title, company, location, salary, and skill normalization.",
            ],
            [
                "Deduplication",
                "Use duplicate_hash and content_hash to detect cross-source duplicates and changed postings.",
            ],
            [
                "Gold analytics",
                "Start with monthly skill counts, source coverage, active postings, and salary summary datasets.",
            ],
        ],
        [1.7, 4.8],
    )

    add_heading(doc, "Appendix A. Phase 1 Design Principles Preserved", 1)
    for item in [
        "Source-agnostic architecture: every source plugs into a common extractor interface.",
        "Raw data preservation: bronze records keep original payloads and do not overwrite existing keys.",
        "Incremental processing: extractor checkpoint hooks and content hashes prepare for resumable crawls.",
        "Configuration over code: sources, skills, titles, salary, locations, logging, and paths live in YAML.",
        "Analytics-first design: canonical fields support measurable questions about skills, salary, geography, companies, experience, and education.",
        "Extensibility: storage, extraction, normalization, enrichment, analytics, and dashboard modules can evolve independently.",
    ]:
        add_bullet(doc, item)

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT_PATH)


if __name__ == "__main__":
    build_document()
    print(OUTPUT_PATH)
