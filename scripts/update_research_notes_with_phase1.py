"""Append the completed Phase 1 implementation update to the research notes DOCX."""

# ruff: noqa: E501

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

PROJECT_ROOT = Path(__file__).resolve().parents[1]
SOURCE_DOCX = Path.home() / "Desktop" / "Job Market Intelligence Research Notes.docx"
OUTPUT_DOCX = PROJECT_ROOT / "docs" / "Job Market Intelligence Research Notes - Updated.docx"

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(31, 41, 55)
MUTED = RGBColor(88, 96, 105)
LIGHT_FILL = "F2F4F7"
BORDER = "D9DEE7"


def set_cell_shading(cell, fill: str) -> None:  # noqa: ANN001
    """Apply a background fill to a table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell) -> None:  # noqa: ANN001
    """Set comfortable cell margins."""
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin_name, value in {"top": 80, "bottom": 80, "start": 120, "end": 120}.items():
        node = tc_mar.find(qn(f"w:{margin_name}"))
        if node is None:
            node = OxmlElement(f"w:{margin_name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table) -> None:  # noqa: ANN001
    """Apply a quiet single-grid border."""
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), BORDER)


def set_table_geometry(table, widths: list[float]) -> None:  # noqa: ANN001
    """Set fixed table widths and cell geometry."""
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    table.allow_autofit = False

    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), "9360")
    tbl_w.set(qn("w:type"), "dxa")

    for row in table.rows:
        for index, width in enumerate(widths):
            if index >= len(row.cells):
                continue
            row.cells[index].width = Inches(width)
            row.cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(row.cells[index])
    set_table_borders(table)


def style_run(run, *, bold: bool = False, color: RGBColor | None = None) -> None:  # noqa: ANN001
    """Apply the document's update-section run style."""
    run.font.name = "Calibri"
    run.font.size = Pt(11)
    run.font.bold = bold
    run.font.color.rgb = color or INK


def add_body(doc: Document, text: str) -> None:
    """Add body text to the update section."""
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(text)
    style_run(run)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.1


def add_heading(doc: Document, text: str, level: int) -> None:
    """Add a heading with the notes' restrained blue hierarchy."""
    paragraph = doc.add_heading(text, level=level)
    paragraph.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    paragraph.paragraph_format.space_after = Pt(6)
    for run in paragraph.runs:
        run.font.name = "Calibri"
        run.font.color.rgb = BLUE if level < 3 else DARK_BLUE
        run.font.size = Pt(16 if level == 1 else 13 if level == 2 else 12)


def add_bullet(doc: Document, text: str) -> None:
    """Add a bullet item."""
    paragraph = doc.add_paragraph(style="List Paragraph")
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = num_pr.find(qn("w:ilvl"))
    if ilvl is None:
        ilvl = OxmlElement("w:ilvl")
        num_pr.append(ilvl)
    ilvl.set(qn("w:val"), "0")
    num_id = num_pr.find(qn("w:numId"))
    if num_id is None:
        num_id = OxmlElement("w:numId")
        num_pr.append(num_id)
    num_id.set(qn("w:val"), "14")
    run = paragraph.add_run(text)
    style_run(run)
    paragraph.paragraph_format.space_after = Pt(4)


def add_table(
    doc: Document, headers: list[str], rows: list[list[str]], widths: list[float]
) -> None:
    """Add a compact fixed-width table."""
    table = doc.add_table(rows=1, cols=len(headers))
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        set_cell_shading(cell, LIGHT_FILL)
        run = cell.paragraphs[0].add_run(header)
        style_run(run, bold=True, color=DARK_BLUE)
    for row_values in rows:
        row = table.add_row()
        for index, value in enumerate(row_values):
            run = row.cells[index].paragraphs[0].add_run(value)
            style_run(run)
    set_table_geometry(table, widths)
    doc.add_paragraph()


def add_phase1_update(doc: Document) -> None:
    """Append the implementation update content."""
    doc.add_page_break()
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = title.add_run("Phase 1 Implementation Update")
    style_run(run, bold=True, color=DARK_BLUE)
    run.font.size = Pt(20)
    title.paragraph_format.space_after = Pt(4)

    subtitle = doc.add_paragraph()
    run = subtitle.add_run("Updated with repository implementation details as of August 17, 2026.")
    style_run(run, color=MUTED)
    subtitle.paragraph_format.space_after = Pt(10)

    add_body(
        doc,
        "Phase 1 is complete. The project now has a tested, source-agnostic foundation for "
        "future data acquisition, processing, enrichment, analytics, and dashboard work. The "
        "completed foundation includes configuration, logging, canonical data models, extractor "
        "contracts, storage layers, checkpointing, retry behavior, hashing utilities, and tests.",
    )

    add_heading(doc, "Phase 1 Completion Summary", 1)
    add_table(
        doc,
        ["Step", "Actual implementation", "Status"],
        [
            [
                "Project structure",
                "Modular Python layout for config, models, extractors, storage, processors, utilities, scripts, tests, analytics, enrichment, and dashboard placeholders.",
                "Completed",
            ],
            [
                "Configuration",
                "Singleton Config loader validates YAML with Pydantic, supports .env loading, and includes selected JMIP_* environment overrides.",
                "Completed",
            ],
            [
                "Canonical model",
                "Strict Pydantic JobPosting schema with source, company, position, location, salary, employment, experience, education, description, skills, raw, and metadata sections.",
                "Completed",
            ],
            [
                "Extractor interface",
                "BaseExtractor provides search/fetch/parse hooks plus extract() orchestration, raw preservation, checkpointing, result summaries, and source config access.",
                "Completed",
            ],
            [
                "Data layers",
                "BaseStorage plus JSON-backed BronzeStorage, SilverStorage, and GoldStorage with safe key handling and layer-specific behavior.",
                "Completed",
            ],
            [
                "Logging",
                "Centralized Loguru setup with context fields, bind_context(), log_event(), console/file sinks, and custom SCRAPER level.",
                "Completed",
            ],
            [
                "Utilities",
                "Checkpoint manager, retry decorator, deterministic job/content/duplicate hashing, and config-backed helper infrastructure.",
                "Completed",
            ],
            [
                "Verification",
                "Foundation tests cover configuration, model validation, extractor workflow, storage behavior, checkpoints, hashes, and logging context.",
                "Completed",
            ],
        ],
        [1.2, 4.1, 1.2],
    )

    add_heading(doc, "Current Canonical JobPosting Model", 1)
    add_body(
        doc,
        "The canonical record emitted by future extractors is JobPosting. Canonical sections forbid "
        "unexpected fields, while RawSourceRecord allows source-specific audit payloads. The model "
        "also supports assignment validation and derived metadata generation.",
    )
    add_table(
        doc,
        ["Section", "Implemented fields and behavior"],
        [
            [
                "PostingMetadata",
                "job_id, content_hash, duplicate_hash, ingestion_batch_id, first_seen_at, last_seen_at, posted_at, expires_at, status, is_active, tags; validates date ordering and status-active consistency.",
            ],
            [
                "SourceInfo",
                "name, source_type, url, external_id, fetched_at; source_type uses SourceType enum.",
            ],
            ["CompanyInfo", "name, normalized_name, website, industry."],
            ["PositionInfo", "title, normalized_title, family, seniority."],
            [
                "LocationInfo",
                "raw, city, region, country, workplace_type, latitude, longitude; validates latitude and longitude bounds.",
            ],
            [
                "SalaryInfo",
                "min_amount, max_amount, currency, pay_period, raw_text; validates salary range and normalizes currency to uppercase.",
            ],
            ["ExperienceInfo", "level, min_years, max_years, raw_text; validates year range."],
            ["EducationInfo", "minimum_degree, preferred_degrees, certifications, raw_text."],
            [
                "DescriptionInfo",
                "raw_text, clean_text, html, responsibilities, qualifications, benefits.",
            ],
            ["SkillInfo", "explicit, inferred, categories."],
            [
                "RawSourceRecord",
                "raw_url, raw_payload, raw_text, plus allowed source-specific extras.",
            ],
        ],
        [1.65, 4.85],
    )
    add_bullet(
        doc,
        "Enums now include SourceType, PostingStatus, EmploymentType, WorkplaceType, PayPeriod, ExperienceLevel, and DegreeLevel.",
    )
    add_bullet(
        doc,
        "JobPosting.with_derived_metadata() fills stable job_id, content_hash, and duplicate_hash from canonical identity fields.",
    )

    add_heading(doc, "Current Configuration Inventory", 1)
    add_table(
        doc,
        ["Config file", "Current purpose"],
        [
            [
                "settings.yaml",
                "Environment, project metadata, data/log paths, storage defaults, scraping retry/checkpoint/timeouts/user agent, and processing strictness.",
            ],
            [
                "sources.yaml",
                "Enabled primary ATS sources: Greenhouse, Lever, Ashby. Deferred sources: Indeed, LinkedIn, Workday, and SmartRecruiters.",
            ],
            [
                "skills.yaml",
                "Skill taxonomy for programming languages, data science, machine learning, data engineering, cloud, and AI.",
            ],
            [
                "locations.yaml",
                "Default country, remote/hybrid keywords, regions, and major markets.",
            ],
            [
                "salary.yaml",
                "Supported currencies, pay periods, annual-hours normalization, default period, and salary bands.",
            ],
            ["titles.yaml", "Role families and seniority keywords for later title normalization."],
            [
                "logging.yaml",
                "Console/file sinks, retention/rotation/compression, and structured component/source/job/batch/run/event formats.",
            ],
        ],
        [1.45, 5.05],
    )

    add_heading(doc, "Extractor Interface and Data Layers", 1)
    add_body(
        doc,
        "Concrete source extractors should implement search_jobs(), fetch_job(), and parse_job(). "
        "They can then reuse BaseExtractor.extract() for search, fetch, optional bronze save, parse, "
        "failure collection, and checkpoint persistence.",
    )
    add_table(
        doc,
        ["Layer", "Implemented behavior", "Research use"],
        [
            [
                "Bronze",
                "Immutable JSON raw payload storage; duplicate keys raise FileExistsError.",
                "Preserves source evidence for audit and reprocessing.",
            ],
            [
                "Silver",
                "JSON JobPosting storage with schema validation on save and load.",
                "Holds canonical normalized postings.",
            ],
            [
                "Gold",
                "JSON analytics payload storage with dataframe-like object normalization.",
                "Holds trend tables and analytic summaries.",
            ],
            [
                "BaseStorage",
                "Safe key-to-path translation with traversal protection and common save/load/exists/delete contract.",
                "Keeps backends consistent.",
            ],
        ],
        [1.0, 3.25, 2.25],
    )

    add_heading(doc, "Logging and Operational Utilities", 1)
    add_bullet(doc, "configure_logging() installs console, application, and error sinks from YAML.")
    add_bullet(
        doc,
        "LogContext defines standard component, source, job_id, ingestion_batch_id, run_id, and event fields.",
    )
    add_bullet(
        doc, "bind_context() returns a logger pre-bound for a module, source, job, batch, or run."
    )
    add_bullet(
        doc,
        "log_event() records structured events at standard Loguru levels or the custom SCRAPER level.",
    )
    add_bullet(
        doc,
        "CheckpointManager stores JSON checkpoints for resumable extraction and processing workflows.",
    )
    add_bullet(doc, "retry() provides reusable exponential backoff.")
    add_bullet(
        doc,
        "Hashing utilities generate stable job_id, content_hash, duplicate_hash, and JSON hashes.",
    )

    add_heading(doc, "Verification Snapshot", 1)
    add_table(
        doc,
        ["Gate", "Current result"],
        [
            ["pytest", "11 tests passed."],
            ["ruff", "All checked files passed lint during the Phase 1 completion run."],
            ["mypy", "No issues found across 19 source files in the core Phase 1 surface."],
            [
                "Coverage focus",
                "Config loading, model validation, extractor workflow, bronze/silver/gold storage, unsafe key rejection, checkpoint round trip, hashes, and logging context.",
            ],
        ],
        [1.35, 5.15],
    )

    add_heading(doc, "Phase 2 Handoff", 1)
    add_body(
        doc,
        "Phase 1 intentionally stops at the foundation. The project is ready for data acquisition "
        "work but does not yet perform live crawling, browser automation, NLP skill extraction, "
        "salary parsing, deduplication, DuckDB marts, scheduled jobs, or dashboard analytics.",
    )
    for item in [
        "Implement Greenhouse first using BaseExtractor and BronzeStorage.",
        "Follow with Lever and Ashby once the first concrete extractor proves the interface.",
        "Extend JobPostingNormalizer for configured title, company, location, salary, and skill normalization.",
        "Use duplicate_hash and content_hash for cross-source deduplication and changed-posting detection.",
        "Start Gold analytics with monthly skill counts, source coverage, active postings, and salary summaries.",
    ]:
        add_bullet(doc, item)


def main() -> None:
    """Open the attached notes, append the update, and save a new copy."""
    if not SOURCE_DOCX.exists():
        raise FileNotFoundError(SOURCE_DOCX)
    doc = Document(SOURCE_DOCX)
    add_phase1_update(doc)
    OUTPUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT_DOCX)
    print(OUTPUT_DOCX)


if __name__ == "__main__":
    main()
